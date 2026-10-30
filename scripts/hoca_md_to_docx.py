"""HOCA_BILGILENDIRME_NOTU.md → .docx (Mermaid şemaları PNG'ye render edilerek).

Neden ayrı bir betik: Word Mermaid çizemez. Bu yüzden md içindeki ```mermaid blokları
küçük bir yorumlayıcıyla (düğüm + kenar) PNG'ye çizilip belgeye resim olarak gömülür.

Kullanım:
    <python> scripts/hoca_md_to_docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
MD = ROOT / "HOCA_BILGILENDIRME_NOTU.md"
DOCX = ROOT / "HOCA_BILGILENDIRME_NOTU.docx"
IMG = ROOT / "artifacts" / "figures" / "hoca_semalar"

NODE_RE = r'([A-Za-z_][A-Za-z0-9_]*)\s*(?:\[\s*"([^"]*)"\s*\]|\{\s*"([^"]*)"\s*\})?'


# --------------------------------------------------------------------------- #
# Mermaid → PNG
# --------------------------------------------------------------------------- #
def parse_mermaid(txt: str):
    lines = [ln.strip() for ln in txt.strip().splitlines() if ln.strip()]
    direction = "TD"
    if lines and lines[0].startswith("flowchart"):
        direction = lines[0].split()[1].upper()
        lines = lines[1:]
    labels: dict[str, str] = {}
    shape: dict[str, str] = {}
    edges: list[tuple[str, str, str]] = []

    def touch(nid, lab, kind="box"):
        if lab is not None and lab != "":
            labels[nid] = lab
            shape[nid] = kind
        elif nid not in labels:
            labels[nid] = nid
            shape.setdefault(nid, kind)

    for ln in lines:
        m = re.match(rf'^{NODE_RE}\s*-->(?:\|\s*"([^"]*)"\s*\|)?\s*{NODE_RE}\s*$', ln)
        if not m:
            continue
        a, alab, adia, elab, b, blab, bdia = m.groups()
        touch(a, alab or (adia and adia) or None, "diamond" if adia else "box")
        touch(b, blab or (bdia and bdia) or None, "diamond" if bdia else "box")
        edges.append((a, b, elab or ""))
    return direction, labels, shape, edges


def _depth(nodes, edges):
    indeg = {n: 0 for n in nodes}
    adj = {n: [] for n in nodes}
    for a, b, _ in edges:
        adj[a].append(b)
        indeg[b] += 1
    depth = {n: 0 for n in nodes}
    changed = True
    for _ in range(len(nodes) + 1):
        changed = False
        for a, b, _ in edges:
            if depth[b] < depth[a] + 1:
                depth[b] = depth[a] + 1
                changed = True
        if not changed:
            break
    return depth, adj


def render_mermaid(txt: str, out_png: Path, title: str | None = None):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Polygon

    direction, labels, shape, edges = parse_mermaid(txt)
    depth, _adj = _depth(labels, edges)
    by_depth: dict[int, list[str]] = {}
    for n in labels:
        by_depth.setdefault(depth[n], []).append(n)

    # ---- düğüm boyutları (etiket uzunluğuna göre)
    size = {}
    for n, lab in labels.items():
        ls = [x for x in lab.replace("<br/>", "\n").replace("<br>", "\n").split("\n")]
        wch = max(len(x) for x in ls) if ls else 4
        nl = max(1, len(ls))
        size[n] = (min(4.2, max(2.2, 0.105 * wch + 0.9)), 0.52 * nl + 0.52)
    max_w = max(s[0] for s in size.values())
    slot = max_w + 0.85
    layer = 2.35

    pos = {}
    for d, ns in by_depth.items():
        ns = sorted(ns)
        for i, n in enumerate(ns):
            off = (i - (len(ns) - 1) / 2.0) * slot
            pos[n] = (d * (max_w + 2.0), -off) if direction == "LR" else (-off, -d * layer)

    xs = [p[0] for p in pos.values()]
    ys = [p[1] for p in pos.values()]
    hw = max(size[n][0] for n in pos) / 2 + 0.35
    hh = max(size[n][1] for n in pos) / 2 + 0.35
    x0, x1 = min(xs) - hw, max(xs) + hw
    y0, y1 = min(ys) - hh, max(ys) + hh
    W, H = x1 - x0, y1 - y0
    fig, ax = plt.subplots(figsize=(min(13.5, max(5.0, W * 0.62)), min(9.5, max(1.8, H * 0.62))))
    ax.set_xlim(x0, x1)
    ax.set_ylim(y0, y1)
    ax.set_aspect("equal")
    ax.axis("off")

    def border(n, toward):
        (cx, cy), (w, h) = pos[n], size[n]
        dx, dy = toward[0] - cx, toward[1] - cy
        if abs(dx) < 1e-9 and abs(dy) < 1e-9:
            return cx, cy
        tx = (w / 2) / abs(dx) if abs(dx) > 1e-9 else 1e9
        ty = (h / 2) / abs(dy) if abs(dy) > 1e-9 else 1e9
        t = min(tx, ty)
        return cx + dx * t, cy + dy * t

    # ---- kenarlar (düğümlerden ÖNCE çizilir ki kutular üstte kalsın)
    for a, b, elab in edges:
        p1, p2 = border(a, pos[b]), border(b, pos[a])
        curved = abs(pos[a][0] - pos[b][0]) > 1e-6 and abs(pos[a][1] - pos[b][1]) > 1e-6
        arr = FancyArrowPatch(p1, p2, arrowstyle="-|>", mutation_scale=11,
                              color="#374151", lw=1.15,
                              connectionstyle=f"arc3,rad={0.18 if curved else 0.0}",
                              shrinkA=0, shrinkB=0, zorder=1)
        ax.add_patch(arr)
        if elab:
            ax.text((p1[0] + p2[0]) / 2, (p1[1] + p2[1]) / 2 + 0.14,
                    elab.replace('"', ""), ha="center", va="bottom", fontsize=6.6,
                    color="#6b7280", style="italic", zorder=3)

    # ---- düğümler
    for n, (cx, cy) in pos.items():
        w, h = size[n]
        lab = labels[n].replace("<br/>", "\n").replace("<br>", "\n")
        if shape.get(n) == "diamond":
            ax.add_patch(Polygon([[cx, cy + h * 0.72], [cx + w / 2, cy],
                                  [cx, cy - h * 0.72], [cx - w / 2, cy]],
                                 closed=True, fc="#fff7ed", ec="#b45309", lw=1.2, zorder=2))
        else:
            ax.add_patch(FancyBboxPatch((cx - w / 2, cy - h / 2), w, h,
                                        boxstyle="round,pad=0.02,rounding_size=0.10",
                                        fc="#eef2ff", ec="#3730a3", lw=1.2, zorder=2))
        ax.text(cx, cy, lab, ha="center", va="center", fontsize=7.2, color="#111827",
                linespacing=1.25, zorder=3)

    if title:
        ax.set_title(title, fontsize=8.5, color="#4b5563", pad=4)
    out_png.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out_png, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_png


# --------------------------------------------------------------------------- #
# Markdown → DOCX
# --------------------------------------------------------------------------- #
def add_runs(par, text: str):
    """**kalın** ve `kod` biçimlerini koruyarak paragrafa metin ekler."""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            r = par.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`"):
            r = par.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = par.style.font.size
        else:
            par.add_run(tok)


def build(phase: str = "all"):
    """phase: 'render' (yalnız PNG; matplotlib gerekir) · 'docx' (yalnız belge; python-docx).

    İki ayrı Python gerekir: venv'de matplotlib var ama python-docx yok; paketli runtime'da
    tersi. Bu yüzden PNG'ler bir kez üretilir, belge onları yeniden kullanır.
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    md = MD.read_text(encoding="utf-8")
    blocks = re.split(r"(```mermaid\n.*?```|```\n.*?```)", md, flags=re.S)
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(1.8)
        # A4 (TR akademik standart) — Letter yerine
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)

    fig_no = 0
    for blk in blocks:
        if not blk.strip():
            continue
        if blk.startswith("```mermaid"):
            fig_no += 1
            body = blk[len("```mermaid"):].rstrip("`\n")
            png = IMG / f"sema_{fig_no:02d}.png"
            if phase != "docx":
                render_mermaid(body, png)
            if not png.exists():
                raise SystemExit(f"şema PNG yok: {png} — önce --phase render")
            doc.add_picture(str(png), width=Cm(14.0))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run(f"Şema {fig_no}")
            r.italic = True
            r.font.size = Pt(8.5)
            r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            continue
        if blk.startswith("```"):
            body = blk.strip("`\n")
            for ln in body.splitlines():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.4)
                r = p.add_run(ln if ln else " ")
                r.font.name = "Consolas"
                r.font.size = Pt(8.5)
            continue

        lines = blk.splitlines()
        i = 0
        while i < len(lines):
            ln = lines[i]
            s = ln.strip()
            if not s or set(s) == {"-"}:
                i += 1
                continue
            if s.startswith("#"):
                lvl = len(s) - len(s.lstrip("#"))
                txt = s[lvl:].strip()
                if lvl == 1:
                    doc.add_heading(txt, level=0)
                else:
                    doc.add_heading(txt, level=min(4, lvl - 1))
                i += 1
                continue
            # tablo
            if s.startswith("|") and i + 1 < len(lines) and \
                    set(lines[i + 1].strip().replace("|", "").replace(":", "")) <= {"-", " "}:
                rows = []
                j = i
                while j < len(lines) and lines[j].strip().startswith("|"):
                    rows.append([c.strip() for c in lines[j].strip().strip("|").split("|")])
                    j += 1
                header, body = rows[0], rows[2:]
                t = doc.add_table(rows=1, cols=len(header))
                t.style = "Table Grid"
                for k, cell in enumerate(header):
                    t.rows[0].cells[k].text = ""
                    pr = t.rows[0].cells[k].paragraphs[0]
                    r = pr.add_run(re.sub(r"\*\*", "", cell))
                    r.bold = True
                    r.font.size = Pt(8.5)
                t.rows[0].cells[0].paragraphs[0].runs[0].font.size = Pt(8.5)
                for row in body:
                    cells = t.add_row().cells
                    for k, cell in enumerate(row[:len(header)]):
                        cells[k].text = ""
                        pr = cells[k].paragraphs[0]
                        add_runs(pr, cell)
                        for rr in pr.runs:
                            rr.font.size = Pt(8)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
                i = j
                continue
            # alıntı
            if s.startswith(">"):
                buf = []
                while i < len(lines) and lines[i].strip().startswith(">"):
                    buf.append(lines[i].strip().lstrip(">").strip())
                    i += 1
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Cm(0.5)
                p.paragraph_format.space_before = Pt(3)
                add_runs(p, " ".join(buf))
                for rr in p.runs:
                    rr.font.size = Pt(9.5)
                    rr.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
                continue
            # madde listesi
            if re.match(r"^(\d+\.|[-*])\s", s):
                txt = re.sub(r"^(\d+\.|[-*])\s+", "", s)
                p = doc.add_paragraph(style="List Bullet")
                add_runs(p, txt)
                i += 1
                continue
            # düz paragraf (ardışık satırları birleştir)
            buf = [s]
            i += 1
            while i < len(lines):
                nx = lines[i].strip()
                if (not nx or nx.startswith(("#", "|", ">")) or
                        re.match(r"^(\d+\.|[-*])\s", nx)):
                    break
                buf.append(nx)
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(buf))

    doc.save(DOCX)
    return DOCX, fig_no


if __name__ == "__main__":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    phase = "all"
    if "--phase" in sys.argv:
        phase = sys.argv[sys.argv.index("--phase") + 1]
    if phase == "render":
        # yalnız şemalar: matplotlib gerektirir (venv)
        md = MD.read_text(encoding="utf-8")
        n = 0
        for blk in re.findall(r"```mermaid\n.*?```", md, flags=re.S):
            n += 1
            body = blk[len("```mermaid"):].rstrip("`\n")
            p = render_mermaid(body, IMG / f"sema_{n:02d}.png")
            print(f"  şema {n}: {p.name}")
        print(f"{n} şema render edildi → {IMG.relative_to(ROOT)}")
    else:
        out, n = build(phase)
        print(f"DOCX: {out.relative_to(ROOT)} · {n} şema gömüldü")
