"""Yüklenen şemalı DOCX'i taban alıp son analiz bulgularını uygular.

Tasarım/şema boyutları yüklenen dosyadan OLDUĞU GİBİ korunur; yalnız metin hücreleri
güncellenir (run düzeyinde, biçim kaybı olmadan).

Kullanım:
    <python> scripts/hoca_docx_patch.py --src <yuklenen.docx> --dst HOCA_BILGILENDIRME_NOTU.docx
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def replace_in_runs(par, old: str, new: str) -> bool:
    """Paragrafın RUN'ları arasında gezinen metin değişimi (biçim korunur)."""
    runs = par.runs
    full = "".join(r.text for r in runs)
    idx = full.find(old)
    if idx < 0:
        return False
    end = idx + len(old)
    spans, pos = [], 0
    for r in runs:
        spans.append((pos, pos + len(r.text)))
        pos += len(r.text)
    s = e = None
    for i, (a, b) in enumerate(spans):
        if s is None and a <= idx < b:
            s = i
        if a < end <= b:
            e = i
    if s is None or e is None:
        return False
    sa, _ = spans[s]
    ea, _ = spans[e]
    prefix = runs[s].text[:idx - sa]
    suffix = runs[e].text[end - ea:]
    if s == e:
        runs[s].text = prefix + new + suffix
    else:
        runs[s].text = prefix + new
        for i in range(s + 1, e):
            runs[i].text = ""
        runs[e].text = suffix
    return True


P39_OLD = "0,932–0,975"
P39_NEW = "0,932–0,941"

P44_OLD = "0,932 (Top-1) / 0,975 (Noisy-OR) hasta AUC değerine ulaştı."
P44_NEW = (
    "0,932 Top-1 hasta AUC değerine ulaştı. Burada dürüst bir not düşmem gerekiyor: JSRT’nin "
    "149 görüntüsünün tamamı nodüllü olduğu için “hasta AUC” ancak ortak bir sağlıklı kontrol "
    "havuzuyla hesaplanabiliyor ve o havuz modelin eğitimde gördüğü bir dağılım. Bu yüzden bu "
    "satırdaki asıl güvenilir kanıt, tek bir kontrol havuzuna bağlı olmayan “tümörün yerini "
    "bulma” başarısıdır: ilk 10 kutuda %85,2, ilk 50 kutuda %94,0. (Kontrol havuzunu JSRT’nin "
    "kendi fiziksel piksel boyutuna daralttığımızda AUC düşmüyor, 0,932 → 0,941 oluyor; yani "
    "sayı havuz seçiminden şişmiyor.)"
)

CELL1_OLD = "0,932 (Top-1) / 0,975 (Noisy-OR) AUC · %94,0 Recall@50"
CELL1_NEW = "İlk 10 kutuda %85,2 · ilk 50 kutuda %94,0 · Top-1 hasta AUC 0,932"

CELL2_OLD = "yüksek başarısını koruyor"
CELL2_NEW = ("yüksek başarısını koruyor. (NODE21’de sağlıklı kontroller ortak havuzdan geldiği "
             "için asıl güvenilir dış ölçüt, tümörün yerini bulma başarısıdır: %94,0.)")


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", required=True)
    ap.add_argument("--dst", default=str(ROOT / "HOCA_BILGILENDIRME_NOTU.docx"))
    args = ap.parse_args()

    from docx import Document

    src, dst = Path(args.src), Path(args.dst)
    if not src.exists():
        raise SystemExit(f"kaynak yok: {src}")
    if src.resolve() != dst.resolve():
        try:
            shutil.copy(src, dst)
        except PermissionError:
            # Word dosyayı açık tutuyor olabilir → yeni sürüm adıyla yaz
            alt = dst.with_name(dst.stem + "_SEMALI_10.docx")
            shutil.copy(src, alt)
            print(f"  ⚠️ {dst.name} kilitli (Word açık) → {alt.name} olarak yazılıyor")
            dst = alt
    doc = Document(dst)

    done = {}
    done["p39"] = replace_in_runs(doc.paragraphs[39], P39_OLD, P39_NEW)
    done["p44"] = replace_in_runs(doc.paragraphs[44], P44_OLD, P44_NEW)

    tbl = doc.tables[0]
    c1 = tbl.rows[6].cells[1].paragraphs[0]
    c2 = tbl.rows[6].cells[2].paragraphs[0]
    done["tablo_JSRT_sutun2"] = replace_in_runs(c1, CELL1_OLD, CELL1_NEW)
    done["tablo_JSRT_sutun3"] = replace_in_runs(c2, CELL2_OLD, CELL2_NEW)

    for k, v in done.items():
        print(f"  {'✔' if v else '✘ DEĞİŞMEDİ'} {k}")
    if not all(done.values()):
        raise SystemExit("bazı değişiklikler uygulanamadı — hedef metin bulunamadı")

    doc.save(dst)
    print(f"\n  yazıldı: {dst}  ({dst.stat().st_size / 1e6:.2f} MB)")


if __name__ == "__main__":
    main()
